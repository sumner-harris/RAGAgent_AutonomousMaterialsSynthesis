from pathlib import Path

import streamlit as st
from docx import Document as DocxDocument

from RAG.generation.model_selector import (
    load_prompt_model_options,
    load_rerank_model_options,
)
from RAG.generation.streaming import stream_answer
from RAG.retrieval.retriever import QAContextRetriever
from ingestion.loaders import process_uploads_for_session
from state.config import (
    DEFAULT_RETRIEVAL_DIVERSITY,
    DEFAULT_RETRIEVAL_METHOD,
    DEFAULT_RERANK_TOP_K,
    LEGACY_RETRIEVAL_METHOD_ALIASES,
    RETRIEVAL_METHODS,
    VECTOR_RETRIEVAL_METHODS,
)
from RAG.retrieval.reranker import reranking_is_configured


def model_selector(options):
    """Select the generation model for OpenAI or custom endpoints."""
    current = st.session_state.get("gpt_model")
    prompt_provider = st.session_state.get("prompt_provider")
    cache_ready = st.session_state.get("available_prompt_models_cache_key") is not None

    if prompt_provider == "custom" and cache_ready:
        discovered_models = list(options)
        if len(discovered_models) == 1:
            model_id = discovered_models[0]
            st.session_state.custom_gpt_model = model_id
            st.session_state.gpt_model = model_id
            st.text_input(
                "Select model:",
                value=model_id,
                disabled=True,
                help="Discovered automatically from the custom prompt base URL.",
            )
            return

        if discovered_models:
            selector_options = discovered_models + ["Custom model..."]
            if current in discovered_models:
                index = selector_options.index(current)
            elif not current:
                index = 0
            else:
                index = len(selector_options) - 1
            selected = st.selectbox(
                "Select model:",
                selector_options,
                index=index,
            )

            if selected == "Custom model...":
                custom_default = st.session_state.get("custom_gpt_model") or (
                    current if current and current not in discovered_models else ""
                )
                custom_model = st.text_input(
                    "Custom model ID",
                    value=custom_default,
                    help=(
                        "Use this when your endpoint serves a model that does not "
                        "appear in the discovered list."
                    ),
                ).strip()
                st.session_state.custom_gpt_model = custom_model
                st.session_state.gpt_model = custom_model
            else:
                st.session_state.custom_gpt_model = ""
                st.session_state.gpt_model = selected
            return

        custom_model = st.text_input(
            "Custom model ID",
            value=st.session_state.get("custom_gpt_model") or current or "",
            help=(
                "Enter the served model ID manually when your custom endpoint does "
                "not support /v1/models discovery."
            ),
        ).strip()
        st.session_state.custom_gpt_model = custom_model
        st.session_state.gpt_model = custom_model
        return

    selector_options = list(options) + ["Custom model..."]
    if current in options:
        index = selector_options.index(current)
    elif not current and options:
        index = 0
    else:
        index = len(selector_options) - 1
    selected = st.selectbox(
        "Select model:",
        selector_options,
        index=index,
    )

    if selected == "Custom model...":
        custom_default = st.session_state.get("custom_gpt_model") or (
            current if current and current not in options else ""
        )
        custom_model = st.text_input(
            "Custom model ID",
            value=custom_default,
            help="Use this when your endpoint serves a model that does not appear in the discovered list.",
        ).strip()
        st.session_state.custom_gpt_model = custom_model
        st.session_state.gpt_model = custom_model
    else:
        st.session_state.custom_gpt_model = ""
        st.session_state.gpt_model = selected


def rerank_model_selector():
    """Select the rerank model for the configured rerank endpoint."""
    rerank_base_url = st.session_state.get("rerank_api_base_url")
    if not rerank_base_url:
        st.session_state.custom_rerank_model = ""
        st.session_state.rerank_model = ""
        return

    current = st.session_state.get("rerank_model")
    options = load_rerank_model_options(
        api_key=st.session_state.get("api_key"),
        base_url=rerank_base_url,
        session=st.session_state,
    )
    cache_ready = st.session_state.get("available_rerank_models_cache_key") is not None

    if cache_ready:
        discovered_models = list(options)
        if len(discovered_models) == 1:
            model_id = discovered_models[0]
            st.session_state.custom_rerank_model = model_id
            st.session_state.rerank_model = model_id
            st.text_input(
                "Rerank model",
                value=model_id,
                disabled=True,
                help="Discovered automatically from the rerank base URL.",
            )
            return

        if discovered_models:
            selector_options = discovered_models + ["Custom model..."]
            if current in discovered_models:
                index = selector_options.index(current)
            elif not current:
                index = 0
            else:
                index = len(selector_options) - 1
            selected = st.selectbox(
                "Rerank model",
                selector_options,
                index=index,
                help=(
                    "Optional post-retrieval reranker applied after BM25, FAISS, "
                    "or hybrid retrieval."
                ),
            )

            if selected == "Custom model...":
                custom_default = st.session_state.get("custom_rerank_model") or (
                    current if current and current not in discovered_models else ""
                )
                custom_model = st.text_input(
                    "Custom rerank model ID",
                    value=custom_default,
                    help=(
                        "Use this when your rerank endpoint serves a model that "
                        "does not appear in the discovered list."
                    ),
                ).strip()
                st.session_state.custom_rerank_model = custom_model
                st.session_state.rerank_model = custom_model
            else:
                st.session_state.custom_rerank_model = ""
                st.session_state.rerank_model = selected
            return

        custom_model = st.text_input(
            "Rerank model ID",
            value=st.session_state.get("custom_rerank_model") or current or "",
            help=(
                "Enter the served rerank model ID manually when the rerank "
                "endpoint does not support /v1/models discovery."
            ),
        ).strip()
        st.session_state.custom_rerank_model = custom_model
        st.session_state.rerank_model = custom_model

    if reranking_is_configured(
        model=st.session_state.get("rerank_model"),
        base_url=st.session_state.get("rerank_api_base_url"),
    ):
        st.number_input(
            "Keep top reranked chunks",
            min_value=1,
            step=1,
            key="rerank_top_k",
            help=(
                "After reranking, only the highest-scoring chunks up to this "
                "count are sent to answer generation."
            ),
        )


def retrieval_method_selector():
    current = st.session_state.get("retrieval_method", DEFAULT_RETRIEVAL_METHOD)
    current = LEGACY_RETRIEVAL_METHOD_ALIASES.get(current, current)
    options = list(RETRIEVAL_METHODS.keys())
    if current not in options:
        current = DEFAULT_RETRIEVAL_METHOD
    index = options.index(current)
    st.session_state.retrieval_method = st.selectbox(
        "Select retrieval method:",
        options,
        index=index,
        format_func=lambda method: RETRIEVAL_METHODS[method]["label"],
        help="Only the selected retrieval method runs when you click Answer.",
    )
    st.caption(RETRIEVAL_METHODS[st.session_state.retrieval_method]["description"])
    if reranking_is_configured(
        model=st.session_state.get("rerank_model"),
        base_url=st.session_state.get("rerank_api_base_url"),
    ):
        st.caption(
            f"Reranking active via {st.session_state.get('rerank_model')}."
        )


def advanced_controls():
    """Retrieval tuning controls."""
    with st.expander("Advanced Controls"):
        if st.session_state.retrieval_method in VECTOR_RETRIEVAL_METHODS:
            st.slider(
                "Diversity",
                0.0,
                1.0,
                DEFAULT_RETRIEVAL_DIVERSITY,
                0.01,
                key="diversity",
            )
        else:
            st.session_state.diversity = DEFAULT_RETRIEVAL_DIVERSITY
            st.caption("Diversity is only used by FAISS/MMR-based retrieval methods.")


def render_rerank_results(rerank_chunks: list[dict]):
    """Render reranked chunks and their endpoint scores."""
    if not rerank_chunks:
        return

    with st.expander("Rerank Results"):
        st.caption("Chunks kept after reranking, shown in the final reranked order.")
        for chunk in rerank_chunks:
            score = chunk.get("score")
            score_text = f"{score:.4f}" if isinstance(score, (int, float)) else str(score)
            st.markdown(
                f"**#{chunk.get('rank', '?')} | {chunk.get('source', '')} | "
                f"chunk {chunk.get('chunk_id', '?')}**"
            )
            st.caption(f"Rerank score: {score_text}")
            st.markdown(chunk.get("text", "").replace("\n", "<br>"), unsafe_allow_html=True)
            st.divider()


def _render_provenance_evidence(detail: dict) -> None:
    """Render chunk and GraphRAG evidence for one source document."""
    chunk_ids = detail.get("chunk_ids") or []
    if chunk_ids:
        chunk_text = ", ".join(str(chunk_id) for chunk_id in chunk_ids)
        st.caption(f"Chunks used: {chunk_text}")

    graph_lines = []
    if detail.get("graph_source_ids"):
        graph_lines.append(
            "Graph local source rows: "
            + ", ".join(detail["graph_source_ids"])
        )
    if detail.get("graph_report_ids"):
        graph_lines.append(
            "Graph community reports: "
            + ", ".join(detail["graph_report_ids"])
        )
    if detail.get("graph_entity_ids"):
        graph_lines.append(
            "Graph entities: "
            + ", ".join(detail["graph_entity_ids"])
        )
    if detail.get("graph_relationship_ids"):
        graph_lines.append(
            "Graph relationships: "
            + ", ".join(detail["graph_relationship_ids"])
        )
    if detail.get("graph_claim_ids"):
        graph_lines.append(
            "Graph claims: "
            + ", ".join(detail["graph_claim_ids"])
        )
    if graph_lines:
        for line in graph_lines:
            st.caption(line)


def render_source_documents(
    source_document_details: list[dict],
    source_documents: list[str] | None = None,
):
    """Render final source documents kept for answer generation."""
    source_documents = source_documents or []
    if not source_document_details and not source_documents:
        return

    with st.expander("Important Source Documents"):
        st.caption(
            "Documents represented in the final answer context. For GraphRAG, "
            "this is best-effort provenance inferred from GraphRAG citations "
            "and context records."
        )
        if not source_document_details:
            for source in source_documents:
                st.markdown(f"**{source}**")
            return

        for detail in source_document_details:
            evidence_count = detail.get("evidence_count")
            evidence_suffix = (
                f" ({evidence_count} evidence link{'s' if evidence_count != 1 else ''})"
                if isinstance(evidence_count, int) and evidence_count > 0
                else ""
            )
            st.markdown(f"**{detail.get('source', '')}**{evidence_suffix}")
            _render_provenance_evidence(detail)
            st.divider()


def qa_panel(client):
    st.header("Ask a Question")

    has_connection = st.session_state.get("api_verified", False)
    has_kb = st.session_state.get("index") is not None
    disabled = not (has_connection and has_kb)

    if not has_connection:
        st.info("Please configure the endpoint routing above and click Connect.")
    elif not has_kb:
        st.info("Please build or load a Knowledge-Base.")

    st.session_state.setdefault("gpt_model", "gpt-4.1-2025-04-14")
    st.session_state.setdefault("last_query", "")
    st.session_state.setdefault("last_answer", "")
    st.session_state.setdefault("context_meta", "")
    st.session_state.setdefault("rerank_meta", [])
    st.session_state.setdefault("source_documents_meta", [])
    st.session_state.setdefault("source_document_details_meta", [])
    st.session_state.setdefault("rerank_top_k", DEFAULT_RERANK_TOP_K)
    st.session_state.setdefault("retrieval_method", DEFAULT_RETRIEVAL_METHOD)

    model_selector(load_prompt_model_options(client))
    rerank_model_selector()
    retrieval_method_selector()
    advanced_controls()

    query = st.text_area(
        "Ask your question:",
        height=280,
        disabled=disabled,
    )

    st.markdown("#### Add any relevant file(s) for this question (optional)")
    uploaded_files = st.file_uploader(
        "Upload PDFs/TXT/CSV/XLSX or images:",
        type=[
            "pdf",
            "txt",
            "csv",
            "xlsx",
            "png",
            "jpg",
            "jpeg",
            "gif",
            "bmp",
            "tif",
            "tiff",
        ],
        accept_multiple_files=True,
        disabled=disabled,
    )

    col1, _, col3 = st.columns([1, 2, 1])
    answer_clicked = col1.button(
        "Answer", use_container_width=True, disabled=disabled
    )
    save_clicked = col3.button(
        "Save last Q&A", use_container_width=True, disabled=disabled
    )

    retriever = QAContextRetriever()

    class _GraphCallbacks:
        def info(self, msg):
            st.info(msg)

        def warning(self, msg):
            st.warning(msg)

    if answer_clicked and query:
        if not st.session_state.get("gpt_model"):
            st.error("Enter a model ID before asking a question.")
            return

        with st.spinner("Processing uploaded files..."):
            try:
                n_text, n_imgs = process_uploads_for_session(
                    uploaded_files=uploaded_files,
                    client=client,
                    embedding_model=st.session_state.embedding_model,
                    embedding_profile=st.session_state.get("embedding_profile"),
                    dimension=st.session_state.get("dimension"),
                    api_key=st.session_state.api_key,
                    base_url=st.session_state.embedding_api_base_url,
                )
                if n_text or n_imgs:
                    st.success(
                        f"Processed {n_text} text chunks and {n_imgs} image(s) from the uploaded files."
                    )
            except Exception as exc:
                st.session_state.upload_db = None
                st.session_state.upload_meta = []
                st.session_state.upload_images = []
                st.error(f"Upload processing failed: {exc}")

            if st.session_state.get("upload_meta") or st.session_state.get("upload_images"):
                st.caption(
                    f"Uploads ready: {len(st.session_state.get('upload_meta', []))} text chunks, "
                    f"{len(st.session_state.get('upload_images', []))} images."
                )

        with st.spinner("Retrieving context..."):
            retrieval_result = retriever.retrieve(
                query,
                use_uploads=True,
                retrieval_method=st.session_state.retrieval_method,
                callbacks=_GraphCallbacks(),
            )
        st.markdown("### Answer")

        answer = stream_answer(
            client,
            query,
            retrieval_result["context_text"],
            images=st.session_state.get("upload_images"),
        )

        st.session_state.update(
            last_query=query,
            last_answer=answer,
            context_meta=retrieval_result.get(
                "retrieved_context_text",
                retrieval_result["context_text"],
            ),
            rerank_meta=retrieval_result["rerank_chunks"],
            source_documents_meta=retrieval_result.get("source_documents", []),
            source_document_details_meta=retrieval_result.get(
                "source_document_details",
                [],
            ),
        )

    if save_clicked and st.session_state.last_answer:
        path = Path("outputs/qa_pairs.docx")
        doc = DocxDocument(path) if path.exists() else DocxDocument()
        doc.add_heading("Question:", 2)
        doc.add_paragraph(st.session_state.last_query)
        doc.add_heading("Answer:", 2)
        doc.add_paragraph(st.session_state.last_answer)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        st.success(f"Saved to {path}")

    if st.session_state.context_meta:
        st.divider()
        with st.expander("Retrieved Context"):
            st.markdown(
                st.session_state.context_meta.replace("\n", "<br>"),
                unsafe_allow_html=True,
            )
        render_source_documents(
            st.session_state.get("source_document_details_meta", []),
            st.session_state.get("source_documents_meta", []),
        )
        render_rerank_results(st.session_state.get("rerank_meta", []))
